"""아웃바운드 시트(CSV/XLSX) 업로드 → 노션 투두 생성."""
from __future__ import annotations

import csv
import io

import notion_client

# 시트에서 계정명으로 쓸 후보 열 이름 (첫 매칭 사용)
ACCOUNT_COLUMNS = ["참가업체명", "Account", "업체명", "회사명"]
# 본문으로 옮길 부가 정보 열
CONTEXT_COLUMNS = ["부스번호", "Description", "Website", "Industry", "Booth Visit Result"]
OWNER_COLUMN = "Owner"


def _read_rows(filename: str, data: bytes) -> list[dict]:
    if filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        header = [str(c or "").strip() for c in next(rows_iter)]
        return [
            {header[i]: str(cell or "").strip() for i, cell in enumerate(row) if i < len(header)}
            for row in rows_iter
        ]
    # CSV — 구글시트 내보내기는 UTF-8(BOM 포함 가능)
    text = data.decode("utf-8-sig")
    return [dict(r) for r in csv.DictReader(io.StringIO(text))]


def import_sheet(filename: str, data: bytes, campaign: str, only_owner: str | None = None) -> dict:
    """시트를 읽어 노션에 Todo 행을 만든다. only_owner가 주어지면 그 Owner 행만.

    반환: {"created": [...], "skipped_dup": [...], "skipped_owner": n, "skipped_noname": n}
    """
    rows = _read_rows(filename, data)
    if not rows:
        return {"created": [], "skipped_dup": [], "skipped_owner": 0, "skipped_noname": 0}

    account_col = next((c for c in ACCOUNT_COLUMNS if c in rows[0]), None)
    if not account_col:
        raise ValueError(f"계정명 열을 못 찾았어요. 필요한 열: {ACCOUNT_COLUMNS} / 실제 열: {list(rows[0])}")

    existing = {r["account"].strip() for r in notion_client.query_calls()}
    created, skipped_dup = [], []
    skipped_owner = skipped_noname = 0

    for row in rows:
        account = row.get(account_col, "").strip()
        if not account:
            skipped_noname += 1
            continue
        if only_owner and row.get(OWNER_COLUMN, "").strip() != only_owner:
            skipped_owner += 1
            continue
        if account in existing:
            skipped_dup.append(account)
            continue
        body = [f"{col}: {row[col]}" for col in CONTEXT_COLUMNS if row.get(col, "").strip()]
        notion_client.create_call(account, campaign=campaign, body_lines=body)
        existing.add(account)
        created.append(account)

    return {
        "created": created,
        "skipped_dup": skipped_dup,
        "skipped_owner": skipped_owner,
        "skipped_noname": skipped_noname,
    }
